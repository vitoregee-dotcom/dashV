#!/usr/bin/env python3
"""Catálogo de peças em PDF (formato Perkins "Parts Book") -> planilha Excel
pronta pra importar em Cadastros > Aplicações > Importar Catálogo > Planilha.

Não usa IA: lê o texto do PDF por regra e traduz com traducoes.py.
Uso:
  pip install pypdf openpyxl cffi
  python3 scripts/catalogo-pdf/extrair.py CATALOGO.pdf --marca PERKINS --saida pecas.xlsx

Regras do formato (validadas no GN65674N_404D_22 / Perkins 404D-22):
- Só entram páginas com o cabeçalho "Item PartNo. References Qty. Description Notes"
  (capa, índice, instruções e desenhos ficam de fora).
- Linha 2 da página = "<Seção> Plate X", linha 3 = código do grupo.
- Linha de item: <item> [flags] <código> [referências] <qtd> <DESCRIÇÃO> [(notas)].
  O PDF às vezes quebra um item em várias linhas -> junta até completar.
- Flag "F" = "faz parte de": o código é do conjunto/kit pai, não da peça
  (ex.: pistão dentro do kit T436364) -> fica só na aba de conferência.
- Linhas sem código (ex.: "5 1 PISTON RING") são componentes não vendidos
  à parte -> descartadas (listadas no relatório).
- Blocos "History" e "Footnotes" ignorados; "ServiceOptions" entra (são
  opções de sobremedida etc., vendáveis).
- Descrição que continua na linha de baixo (ex.: "GASKET - THERMOSTAT" + "HSG")
  é emendada.
Saída:
- Aba "Importar": 1 linha por código (sem os "faz parte de"), colunas
  codigo/descricao/marca/grupo reconhecidas sozinhas pelo importador; as
  extras têm nomes que NÃO disparam o reconhecimento automático (evitar
  "cod", "ref", "desc", "nome", "marca", "grupo", "sub", "linh", "unid",
  "um" no cabeçalho, senão o importador mapeia errado).
- Grupo = seção real da peça (evita "KIT" quando ela também aparece fora dos kits).
- Aba "Catalogo completo": todas as linhas, na ordem do PDF, pra conferência.
"""
import argparse,re,sys,os,json
sys.path.insert(0,os.path.dirname(os.path.abspath(__file__)))
from traducoes import DESC,SEC

FLAGS={'F','R','a.','b.','c.','E','P','J','N','S'}
FORA={'LABEL','TOOL','LONG ENGINE - ASSEMBLY','LEAFLET'}   # não são peças de venda
KITS={'KIT','KITS DE SERVIÇO','KIT DE RETÍFICA'}

def parse(s):
    tk=s.split()
    if not tk or not tk[0].isdigit():return None
    item=int(tk[0]);k=1;part=None;parte_de=False
    while k<len(tk):
        x=tk[k]
        if x=='F':parte_de=True
        if x in FLAGS:k+=1;continue
        if re.search(r'\d',x) and len(x)>=5 and re.fullmatch(r'[A-Z0-9]+',x):part=x;k+=1;break
        return ('semcodigo',None)
    if not part:return None
    for j in range(k,len(tk)-1):
        if tk[j].isdigit() and re.match(r'^[A-Z(]',tk[j+1]):
            desc=' '.join(tk[j+1:])
            notas=' '.join(re.findall(r'\(\d+\)',desc));desc=re.sub(r'\s*(\(\d+\))+\s*$','',desc).strip()
            return ('ok',{'item':item,'parte_de':parte_de,'codigo':part,'qtd':int(tk[j]),'descricao_en':desc,'refs':' '.join(tk[k:j]),'notas':notas})
    return None

def extrair(pdf):
    import pypdf
    r=pypdf.PdfReader(pdf);itens=[];descartes=[];pags=0
    for pi,p in enumerate(r.pages):
        ls=[l.strip() for l in (p.extract_text() or '').split('\n') if l.strip()]
        if not any(l.startswith('Item PartNo.') for l in ls):continue
        pags+=1
        i0=[k for k,l in enumerate(ls) if l.startswith('Item PartNo.')][0]
        cab=ls[1:i0]
        secao=re.sub(r'\s*Plate\s+\S+$','',cab[0]) if cab else ''
        grupo_cod=cab[1] if len(cab)>1 else ''
        modo='itens';ult=None;buf=None
        for l in ls[i0+1:]:
            if l in('History','Footnotes'):modo='fora';buf=None;continue
            if l=='ServiceOptions':modo='itens';continue
            if modo!='itens':continue
            cand=(buf+' '+l) if buf else l
            res=parse(cand)
            if res and res[0]=='ok':
                ult=dict(res[1],pagina=pi+1,secao=secao,grupo_cod=grupo_cod);itens.append(ult);buf=None;continue
            if res and res[0]=='semcodigo':descartes.append((pi+1,cand));buf=None;ult=None;continue
            if buf is None and re.match(r'^\d+\s',l):buf=l;continue
            if buf is not None:
                buf=cand
                if len(buf)>200:descartes.append((pi+1,buf));buf=None
                continue
            if ult and re.fullmatch(r"[A-Z0-9 .,/&'\-]+",l):ult['descricao_en']+=' '+l;continue
            ult=None
    return itens,descartes,pags

def main():
    ap=argparse.ArgumentParser();ap.add_argument('pdf');ap.add_argument('--marca',default='PERKINS');ap.add_argument('--saida',default='pecas.xlsx')
    a=ap.parse_args()
    itens,descartes,pags=extrair(a.pdf)
    itens=[i for i in itens if i['descricao_en'] not in FORA]
    falta_d=sorted(set(i['descricao_en'] for i in itens)-set(DESC));falta_s=sorted(set(i['secao'] for i in itens)-set(SEC))
    tr=lambda d:DESC.get(d,d);ts=lambda s:SEC.get(s,s.upper())
    uniq={}
    for i in itens:
        if i['parte_de']:continue
        u=uniq.setdefault(i['codigo'],dict(i,secoes=[],pags=[]))
        if ts(i['secao']) not in u['secoes']:u['secoes'].append(ts(i['secao']))
        if i['pagina'] not in u['pags']:u['pags'].append(i['pagina'])
    grupo=lambda u:([x for x in u['secoes'] if x not in KITS] or u['secoes'])[0]
    import openpyxl
    from openpyxl.styles import Font,PatternFill,Alignment
    wb=openpyxl.Workbook();ws=wb.active;ws.title='Importar'
    ws.append(['codigo','descricao','marca','grupo','texto em ingles','qtd no motor','aparece em','pagina'])
    for c,u in uniq.items():ws.append([c,tr(u['descricao_en']),a.marca,grupo(u),u['descricao_en'],u['qtd'],', '.join(u['secoes']),', '.join(map(str,u['pags']))])
    h=Font(bold=True,color='FFFFFF');f=PatternFill('solid',fgColor='1C2940')
    for cell in ws[1]:cell.font=h;cell.fill=f;cell.alignment=Alignment(vertical='center')
    for col,w in zip('ABCDEFGH',[14,42,11,30,34,12,40,12]):ws.column_dimensions[col].width=w
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    ws2=wb.create_sheet('Catalogo completo')
    ws2.append(['pagina','grupo','cod. grupo','item','codigo','qtd','descricao','texto em ingles','obs catalogo'])
    for i in itens:ws2.append([i['pagina'],ts(i['secao']),i['grupo_cod'],i['item'],i['codigo'],i['qtd'],tr(i['descricao_en'])+(' (FAZ PARTE DO '+i['codigo']+')' if i['parte_de'] else ''),i['descricao_en'],(i['refs']+' '+i['notas']).strip()])
    for cell in ws2[1]:cell.font=h;cell.fill=f
    for col,w in zip('ABCDEFGHI',[8,30,12,6,14,6,50,34,18]):ws2.column_dimensions[col].width=w
    ws2.freeze_panes='A2';ws2.auto_filter.ref=ws2.dimensions
    wb.save(a.saida)
    print(f'{a.saida}: {pags} paginas de lista, {len(itens)} linhas, {len(uniq)} pecas unicas, {sum(1 for i in itens if i["parte_de"])} "faz parte de"')
    print(f'descartadas (sem codigo): {len(descartes)}');[print('   p.',p,'|',d) for p,d in descartes[:20]]
    if falta_d or falta_s:
        print('\nFALTA TRADUZIR (acrescente em traducoes.py e rode de novo):')
        [print('  DESC',repr(d)) for d in falta_d];[print('  SEC ',repr(s)) for s in falta_s]
if __name__=='__main__':main()
